"""Susun draf bagian (docs/paper/sec*-draft.md) menjadi satu naskah Word untuk ditinjau.

Hanya prosa naskah yang diambil: dari judul "## N." pertama sampai sebelum catatan
kerja ("## Catatan penyusunan", "## Checklist", "## Audit adversarial"). Bagian yang
belum ditulis ditampilkan sebagai penanda, bukan dikarang.

    python scripts/build_docx.py            -> docs/paper/manuscript-draft.docx
"""

from __future__ import annotations

import datetime
import pathlib
import re
import subprocess
import sys

import pypandoc
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

ROOT = pathlib.Path(__file__).resolve().parents[1]
PAPER = ROOT / "docs" / "paper"
KELUAR = PAPER / "manuscript-draft.docx"

# (nomor bagian, judul, berkas draf atau None bila belum ditulis)
URUTAN = [
    ("1", "Introduction", None),
    ("2", "Related Work", "sec2-draft.md"),
    ("3–4", "Preliminaries and Problem Formulation", "sec3-4-draft.md"),
    ("5", "Methods", "sec5-draft.md"),
    ("6–7", "Datasets and Experimental Setup", "sec6-7-draft.md"),
    ("8", "Results", "sec8-draft.md"),
    ("9", "Ablation", None),
    ("10", "Discussion", None),
    ("11", "Threats to Validity", "sec11-draft.md"),
    ("12", "Conclusion", None),
]
CATATAN_KERJA = re.compile(r"^## (Catatan penyusunan|Checklist|Audit adversarial)", re.M)
BAGIAN_NASKAH = re.compile(r"^## \d+\.", re.M)


def ambil_prosa(teks: str) -> str:
    awal = BAGIAN_NASKAH.search(teks)
    if awal is None:
        raise ValueError("tidak ada judul '## N.' di draf")
    akhir = CATATAN_KERJA.search(teks, awal.start())
    isi = teks[awal.start() : akhir.start() if akhir else len(teks)]
    return re.sub(r"\n---\s*$", "\n", isi.rstrip()) + "\n"


def susun_markdown() -> str:
    commit = subprocess.run(["git", "rev-parse", "--short", "HEAD"], cwd=ROOT,
                            capture_output=True, text=True).stdout.strip() or "?"
    potong = [
        "# Block-Level Feasibility of Conformal Calibration on Clinical ECG Data: An Empirical Audit\n",
        f"*Working draft for internal review — generated {datetime.date.today():%Y-%m-%d} "
        f"from commit `{commit}`. Title and abstract are provisional. "
        "Sections marked NOT YET WRITTEN are intentionally empty.*\n",
    ]
    for nomor, judul, berkas in URUTAN:
        if berkas is None:
            potong.append(f"## {nomor}. {judul}\n\n**[NOT YET WRITTEN]**\n")
        else:
            potong.append(ambil_prosa((PAPER / berkas).read_text(encoding="utf-8")))
    return "\n".join(potong)


def beri_garis_tabel(doc: Document) -> None:
    for tabel in doc.tables:
        tblPr = tabel._tbl.tblPr
        garis = OxmlElement("w:tblBorders")
        for sisi in ("top", "left", "bottom", "right", "insideH", "insideV"):
            e = OxmlElement(f"w:{sisi}")
            e.set(qn("w:val"), "single")
            e.set(qn("w:sz"), "4")
            e.set(qn("w:color"), "808080")
            garis.append(e)
        tblPr.append(garis)


def rapikan(path: pathlib.Path) -> None:
    doc = Document(path)
    for nama in ("Normal", "Body Text", "First Paragraph", "Compact"):
        if nama in [s.name for s in doc.styles]:
            gaya = doc.styles[nama]
            gaya.font.name = "Times New Roman"
            gaya.font.size = Pt(11)
            gaya.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    for p in doc.paragraphs:
        if p.style.name == "Title":
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    beri_garis_tabel(doc)
    doc.save(path)


def main() -> int:
    md = susun_markdown()
    pypandoc.convert_text(
        md, "docx", format="markdown+tex_math_dollars+pipe_tables",
        outputfile=str(KELUAR),
        extra_args=["--shift-heading-level-by=-1", "--toc", "--toc-depth=2"],
    )
    rapikan(KELUAR)
    print(f"Ditulis: {KELUAR.relative_to(ROOT)}  ({KELUAR.stat().st_size / 1024:.0f} KB)")
    return 0


if __name__ == "__main__":
    sys.exit(main())
